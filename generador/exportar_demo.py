"""
Saca las tablas de la pestaña PANEL_DATOS del sheet ya recalculado y las deja en dashboard/demo/
como JSON (una lista de filas por tabla). Son los datos que lleva dentro el dashboard de demostración.

Hace lo mismo que los loaders del dashboard (dashboard/loaders.js) con el .xlsx que exporta Google:
busca «## NOMBRE», toma la fila siguiente como cabeceras y lee hasta la primera fila vacía.

Uso:  python generador/exportar_demo.py      (después de build.py y del recálculo con LibreOffice)
"""
import datetime as dt
import json
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
XLSX = RAIZ / "plantilla" / "Plan_financiero_demo.xlsx"
SALIDA = RAIZ / "dashboard" / "demo"
TABLAS = {  # id del dataset: (tabla en PANEL_DATOS, columnas con fecha)
    "kpi": ("KPI", []), "meses": ("MESES", ["mes"]), "casa": ("CASA", []), "casa_meses": ("CASA_MESES", ["mes"]),
    "proyeccion": ("PROYECCION", []), "jubilacion": ("JUBILACION", []), "hitos": ("HITOS", []),
    "viajes": ("VIAJES", ["fecha"]), "viajes_escenarios": ("VIAJES_ESCENARIOS", []),
}

ws = load_workbook(XLSX, data_only=True)["PANEL_DATOS"]
grid = [[c for c in fila] for fila in ws.iter_rows(values_only=True)]
if all(v is None for f in grid for v in f[1:2]):
    raise SystemExit("El .xlsx no tiene valores calculados: recalcúlalo antes (docs/GUIA_DESPLIEGUE.md)")


def valor(v, es_fecha):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer() and not es_fecha:
        return int(v)
    return v


def tabla(nombre, fechas):
    i = next(k for k, f in enumerate(grid) if f and f[0] == "## " + nombre)
    columnas = [c for c in grid[i + 1] if c not in ("", None)]
    filas = []
    for f in grid[i + 2:]:
        if not f or f[0] in ("", None):
            break
        filas.append({c: valor(f[j] if j < len(f) else None, c in fechas) for j, c in enumerate(columnas)})
    return filas


SALIDA.mkdir(parents=True, exist_ok=True)
for id_, (nombre, fechas) in TABLAS.items():
    filas = tabla(nombre, fechas)
    if id_ == "kpi":
        for r in filas:
            if r["clave"] == "fecha_datos" and isinstance(r["valor"], (int, float)):
                r["valor"] = (dt.date(1899, 12, 30) + dt.timedelta(days=int(r["valor"]))).isoformat()
    if id_ == "proyeccion":
        filas = [r for r in filas if r["edad"] != ""]
    (SALIDA / f"{id_}.json").write_text(json.dumps(filas, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{id_}: {len(filas)} filas")
