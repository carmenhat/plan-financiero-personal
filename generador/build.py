"""
Genera el sheet del plan financiero (.xlsx) con openpyxl a partir de generador/datos-demo/.
Después hay que recalcularlo (LibreOffice) y subirlo a Google: ver docs/GUIA_DESPLIEGUE.md.

Uso:  python generador/build.py      → plantilla/Plan_financiero_demo.xlsx
"""
import datetime, json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import AreaChart, LineChart, BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F = "Arial"
NAVY = "1F3A4D"; SECT = "E8EEF2"; INFILL = "FFF7D6"; GOOD = "2E7D32"
EUR = '#,##0 €;-#,##0 €;"–"'
EUR0 = '#,##0 €'
PCT = '0%'; PCT1 = '0.0%'

wb = Workbook()

def font(**k):
    k.setdefault("name", F); return Font(**k)

def title(ws, text, sub, width_cols):
    ws["A1"] = text
    ws["A1"].font = font(size=16, bold=True, color="FFFFFF")
    for c in range(1, width_cols + 1):
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 30
    ws["A2"] = sub
    ws["A2"].font = font(size=9, italic=True, color="555555")

def section(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text).font = font(size=11, bold=True, color=NAVY)
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=SECT)

def header(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = font(size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def inp(cell, value, fmt=None):
    cell.value = value
    cell.font = font(color="0000FF", bold=True)
    cell.fill = PatternFill("solid", fgColor=INFILL)
    if fmt: cell.number_format = fmt

def link(cell, formula, fmt=None):
    cell.value = formula
    cell.font = font(color="008000")
    if fmt: cell.number_format = fmt

def calc(cell, formula, fmt=None, bold=False, size=10, color="000000"):
    cell.value = formula
    cell.font = font(bold=bold, size=size, color=color)
    if fmt: cell.number_format = fmt

def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

# ---------------------------------------------------------------- sheets
ws_lee = wb.active; ws_lee.title = "LEEME"
ws_fut = wb.create_sheet("MI FUTURO")
ws_dat = wb.create_sheet("MIS DATOS")
ws_mov = wb.create_sheet("MOVIMIENTOS")
ws_flex = wb.create_sheet("FLEXIBLE")
ws_casa = wb.create_sheet("CASA")
ws_dec = wb.create_sheet("DECISIONES")
ws_his = wb.create_sheet("HISTORIAL")
ws_viv = wb.create_sheet("VIVIENDA")
ws_inm = wb.create_sheet("INMUEBLES")
ws_via = wb.create_sheet("VIAJES")
ws_par = wb.create_sheet("PARÁMETROS")
ws_reg = wb.create_sheet("REGLAS")
ws_cat = wb.create_sheet("CATEGORÍAS")
ws_pro = wb.create_sheet("PROYECCIÓN")

LAST = 245
MV = "'MOVIMIENTOS'!"; MVR = lambda c: f"{MV}${c}$4:${c}$20000"
FXR = lambda c: f"FLEXIBLE!${c}$4:${c}$200"
DR = lambda col: f"'MIS DATOS'!${col}$6:${col}${LAST}"
NDAT = f"COUNT('MIS DATOS'!$A$6:$A${LAST})"
AQUI = Path(__file__).resolve().parent
DATOS = AQUI / "datos-demo"
CFG = json.loads((DATOS / "config.json").read_text(encoding="utf-8"))
MOVS = json.loads((DATOS / "movimientos.json").read_text(encoding="utf-8"))
SALDOS = json.loads((DATOS / "saldos.json").read_text(encoding="utf-8"))
CU = CFG["cuentas"]; PER = CFG["persona"]; VIV = CFG["vivienda"]; VJ = CFG["viajes"]
ANIO = PER["anio"]; JUB = PER["jubilacion"]
d_ = lambda iso: datetime.date.fromisoformat(iso) if iso else None
P = "'PARÁMETROS'!"; PR = "'PROYECCIÓN'!"; D = "'MIS DATOS'!"; DE = "DECISIONES!"; PI = "VIVIENDA!"; V = "VIAJES!"; IN = "INMUEBLES!"

# ================================================================ PARÁMETROS
ws = ws_par
title(ws, "PARÁMETROS", "Supuestos del plan. Amarillo/azul: los cambias tú. Verde: salen de tus movimientos (puedes escribir encima una cifra).", 3)
params = [
    ("edad", "Edad actual", PER["edad"], "0", ""),
    ("jub", "Edad de jubilación", JUB, "0", ""),
    ("anio", "Año actual", ANIO, "0", "Año de la primera fila de la proyección"),
    ("r_inv", "Rentabilidad anual de la inversión", 0.05, PCT1, "Supuesto nominal, antes de impuestos"),
    ("r_pen", "Rentabilidad anual del plan de pensiones", 0.04, PCT1, "Supuesto nominal"),
    ("r_ef", "Rentabilidad anual del efectivo", f"=IFERROR(SUMIFS({MVR('E')},{MVR('G')},\"Intereses\",{MVR('A')},\">=\"&EDATE(TODAY(),-12))/SUM(INDEX({DR('B')},{NDAT}),INDEX({DR('C')},{NDAT}),INDEX({DR('D')},{NDAT})),0)", "0.00%", "Intereses de los últimos 12 meses ÷ tu efectivo de hoy. Si tus cuentas no pagan intereses, sale casi cero"),
    ("infl", "Inflación anual", 0.025, PCT1, "Para pasar a € de hoy"),
    ("ap_inv", "Aportación mensual a inversión", 0, EUR0, "Hoy no inviertes: pruébalo en DECISIONES"),
    ("ap_pen", "Aportación mensual al plan de pensiones", 0, EUR0, "Si no tienes plan de pensiones, déjalo en 0"),
    ("ap_ef", "Ahorro mensual que dejas en efectivo", f"=IFERROR(AVERAGEIFS({DR('I')},{DR('N')},1),0)", EUR0, "Media de tus meses completos en MIS DATOS"),
    ("g", "Subida anual de las aportaciones", 0.02, PCT1, "Si tu sueldo sube, aportas algo más"),
    ("gastos", "Gastos mensuales (€ de hoy)", f"=IFERROR(AVERAGEIFS({DR('H')},{DR('N')},1),0)", EUR0, "Media de tus meses completos. Ojo: no incluye lo que queda por revisar, ni lo que pague otra persona o se pague en efectivo. Si sabes que gastas más, escribe tu cifra encima"),
    ("ret", "Regla de retirada", 0.04, PCT1, "Regla del 4 %: lo que podrías sacar al año sin agotar el capital"),
    ("pension", "Pensión pública estimada al mes (€ de hoy, 14 pagas)", PER["pension_mes"], EUR0, "Ficticio. Consulta la tuya en Tu Seguridad Social; se supone que sube con la inflación"),
    ("colchon", "Meses de colchón de emergencia", 6, "0", ""),
    ("rev", "Revalorización anual de la vivienda", 0.03, PCT1, "Supuesto"),
    ("r_jub", "Rentabilidad del capital en la jubilación", 0.03, PCT1, "Más prudente que en la fase de ahorro"),
]
PP = {}
header(ws, 3, ["Parámetro", "Valor", "Nota"])
for i, (key, lab, val, fmt, note) in enumerate(params):
    r = 4 + i
    ws.cell(row=r, column=1, value=lab).font = font()
    if isinstance(val, str) and val.startswith("="): link(ws.cell(row=r, column=2), val, fmt)
    else: inp(ws.cell(row=r, column=2), val, fmt)
    ws.cell(row=r, column=3, value=note).font = font(size=9, color="666666")
    PP[key] = f"{P}$B${r}"
lvl_row = 4 + len(params) + 2
section(ws, lvl_row, "NIVELES (por % de libertad financiera)", 3)
header(ws, lvl_row + 1, ["Desde", "Nivel", ""])
niveles = [(0, "🌱 Semilla"), (0.05, "🛡️ Escudo"), (0.15, "🧭 Brújula"), (0.30, "🏗️ Cimientos"),
           (0.50, "🌉 Puente"), (0.75, "🏔️ Cumbre"), (1.0, "🕊️ Libre")]
n0 = lvl_row + 2
for i, (u, n) in enumerate(niveles):
    inp(ws.cell(row=n0 + i, column=1), u, PCT)
    inp(ws.cell(row=n0 + i, column=2), n)
UMB = f"{P}$A${n0}:$A${n0+len(niveles)-1}"
NIV = f"{P}$B${n0}:$B${n0+len(niveles)-1}"
widths(ws, {"A": 42, "B": 16, "C": 60})
ws.freeze_panes = "A4"

# ================================================================ MOVIMIENTOS
ws = ws_mov
title(ws, "MOVIMIENTOS", "Los trae el importador (menú Plan financiero). Tú solo cambias Categoría y escribes Nota; lo amarillo está por revisar.", 11)
header(ws, 3, ["Fecha", "Mes", "Cuenta", "Concepto", "Importe", "Saldo", "Categoría", "Tipo", "Grupo", "Nota", "Clave"])
dvc = DataValidation(type="list", formula1="'CATEGORÍAS'!$A$4:$A$100", allow_blank=True)
ws.add_data_validation(dvc)
for i, m in enumerate(MOVS):
    r = 4 + i
    y, mo, dd_ = map(int, m["fecha"].split("-"))
    ws.cell(row=r, column=1, value=datetime.date(y, mo, dd_)).number_format = "dd/mm/yyyy"
    calc(ws.cell(row=r, column=2), f"=DATE(YEAR(A{r}),MONTH(A{r}),1)", "mmm yyyy")
    ws.cell(row=r, column=3, value=m["cuenta"])
    ws.cell(row=r, column=4, value=m["concepto"])
    ws.cell(row=r, column=5, value=m["importe"]).number_format = '#,##0.00 €'
    if m["saldo"] != "": ws.cell(row=r, column=6, value=m["saldo"]).number_format = '#,##0.00 €'
    inp(ws.cell(row=r, column=7), m["categoria"])
    calc(ws.cell(row=r, column=8), f"=IFERROR(VLOOKUP(G{r},'CATEGORÍAS'!$A$4:$C$100,2,FALSE),\"Por revisar\")")
    calc(ws.cell(row=r, column=9), f"=IFERROR(VLOOKUP(G{r},'CATEGORÍAS'!$A$4:$C$100,3,FALSE),\"\")")
    inp(ws.cell(row=r, column=10), None)
    ws.cell(row=r, column=11, value=m["clave"]).font = font(size=7, color="AAAAAA")
    for c in (1, 3, 4, 5, 6): ws.cell(row=r, column=c).font = font()
MOV_END = 3 + len(MOVS)
dvc.add(f"G4:G5000")
from openpyxl.formatting.rule import FormulaRule
ws.conditional_formatting.add("A4:J5000", FormulaRule(formula=['$H4="Por revisar"'], fill=PatternFill("solid", fgColor="FFF3B0")))
ws.auto_filter.ref = f"A3:K{MOV_END}"
widths(ws, {"A": 11, "B": 10, "C": 16, "D": 62, "E": 12, "F": 12, "G": 24, "H": 11, "I": 12, "J": 36, "K": 14})
ws.freeze_panes = "A4"

# ================================================================ CATEGORÍAS
ws = ws_cat
title(ws, "CATEGORÍAS", "Tipo decide si cuenta como ingreso, gasto o nada (traspaso). Grupo «Casa» es lo que sale en la pestaña CASA. Puedes añadir filas.", 3)
header(ws, 3, ["Categoría", "Tipo", "Grupo"])
dvt = DataValidation(type="list", formula1='"Ingreso,Gasto,Traspaso,Por revisar"'); ws.add_data_validation(dvt)
for i, (c, t, g_) in enumerate(CFG["categorias"]):
    r = 4 + i
    inp(ws.cell(row=r, column=1), c); inp(ws.cell(row=r, column=2), t); inp(ws.cell(row=r, column=3), g_)
dvt.add("B4:B100")
widths(ws, {"A": 30, "B": 14, "C": 18})
ws.freeze_panes = "A4"

# ================================================================ REGLAS
ws = ws_reg
title(ws, "REGLAS", "Si el concepto contiene el texto → esa categoría. Gana la primera que encaje (de arriba abajo). Sin mayúsculas ni acentos que importen.", 5)
header(ws, 3, ["El concepto contiene", "Solo en la cuenta (opcional)", "Signo (+ / − / vacío)", "Categoría", "Nota"])
dvr = DataValidation(type="list", formula1="'CATEGORÍAS'!$A$4:$A$100", allow_blank=True); ws.add_data_validation(dvr)
for i, (pat, cu, sg, cat, nota) in enumerate(CFG["reglas"]):
    r = 4 + i
    inp(ws.cell(row=r, column=1), pat); inp(ws.cell(row=r, column=2), cu or None)
    inp(ws.cell(row=r, column=3), sg or None); inp(ws.cell(row=r, column=4), cat)
    ws.cell(row=r, column=5, value=nota or None).font = font(size=9, color="666666")
dvr.add("D4:D300")
r = 4 + len(CFG["reglas"]) + 1
ws.cell(row=r, column=1, value="Lo que no encaja en ninguna regla (Bizum, ingresos en efectivo…) entra como «❓ Por revisar». Tras añadir reglas: menú Plan financiero → Reaplicar reglas a lo pendiente.").font = font(size=9, color="555555")
widths(ws, {"A": 36, "B": 24, "C": 18, "D": 26, "E": 60})
ws.freeze_panes = "A4"

# ================================================================ FLEXIBLE
ws = ws_flex
title(ws, "FLEXIBLE · retribución flexible", "Lo que gastas con la tarjeta de retribución flexible de tu empresa y se descuenta de la nómina.", 6)
ws["A3"] = "Cuenta a la vez como ingreso y como gasto (Comida y Transporte): tu ahorro no cambia, pero tus gastos reales sí."
ws["A3"].font = font(size=9, color="666666")
header(ws, 3, ["Mes de la nómina", "Comida", "Transporte", "Total", "Nota", "mes (calc)"])
ws["A2"].value = "Lo que gastas con la tarjeta de retribución flexible de tu empresa y se descuenta de la nómina. Copia cada mes esas líneas de tu nómina. Cuenta a la vez como ingreso y como gasto: tu ahorro no cambia, tus gastos reales sí."
ws["A3"].value = "Mes de la nómina"
for i in range(197):
    r = 4 + i
    if i < len(CFG["flexible"]):
        mes_, com_, tra_, nota_ = CFG["flexible"][i]
        inp(ws.cell(row=r, column=1), d_(mes_), "mmm yyyy"); inp(ws.cell(row=r, column=2), com_, '#,##0.00 €')
        inp(ws.cell(row=r, column=3), tra_, '#,##0.00 €'); inp(ws.cell(row=r, column=5), nota_)
    else:
        inp(ws.cell(row=r, column=1), None, "mmm yyyy"); inp(ws.cell(row=r, column=2), None, '#,##0.00 €')
        inp(ws.cell(row=r, column=3), None, '#,##0.00 €'); inp(ws.cell(row=r, column=5), None)
    calc(ws.cell(row=r, column=4), f'=IF(A{r}="","",N(B{r})+N(C{r}))', '#,##0.00 €')
    calc(ws.cell(row=r, column=6), f'=IF(A{r}="","",DATE(YEAR(A{r}),MONTH(A{r}),1))', "mmm yyyy")
    ws.cell(row=r, column=6).font = font(size=8, color="AAAAAA")
widths(ws, {"A": 18, "B": 12, "C": 12, "D": 12, "E": 44, "F": 11})
ws.freeze_panes = "A4"

# ================================================================ MIS DATOS
ws = ws_dat
title(ws, "MIS DATOS", "Se rellena sola: saldos a fin de mes que trae el importador (B–D) e ingresos y gastos de MOVIMIENTOS. Inversión y plan de pensiones (E–F), a mano si algún día los tienes.", 14)
ws["A3"] = "La primera fila es el primer mes del que hay datos de las tres cuentas. Si descargas más histórico de todas, empieza antes."
ws["A3"].font = font(size=9, color="666666")
header(ws, 5, ["Mes", *CU, "Inversión", "Plan de pensiones", "Ingresos del mes", "Gastos del mes",
               "Ahorro del mes", "% de ahorro", "Patrimonio líquido", "Año", "Por revisar", "Mes completo"])
for i, mes in enumerate(SALDOS["meses"]):
    r = 6 + i
    y, mo, _ = map(int, mes.split("-"))
    ws.cell(row=r, column=1, value=datetime.date(y, mo, 1)).number_format = "mmm yyyy"
    for c, cu in zip((2, 3, 4), CU):
        v = SALDOS["saldos"][cu][i]
        ws.cell(row=r, column=c, value=None if v == "" else v).number_format = EUR0
    inp(ws.cell(row=r, column=5), 0, EUR0); inp(ws.cell(row=r, column=6), 0, EUR0)
for r in range(6, LAST + 1):
    calc(ws.cell(row=r, column=7), f'=IF(A{r}="","",SUMIFS({MVR("E")},{MVR("B")},A{r},{MVR("H")},"Ingreso")+SUMIFS({FXR("D")},{FXR("F")},A{r}))', EUR)
    calc(ws.cell(row=r, column=8), f'=IF(A{r}="","",-SUMIFS({MVR("E")},{MVR("B")},A{r},{MVR("H")},"Gasto")+SUMIFS({FXR("D")},{FXR("F")},A{r}))', EUR)
    calc(ws.cell(row=r, column=9), f'=IF(A{r}="","",G{r}-H{r})', EUR)
    calc(ws.cell(row=r, column=10), f'=IF(A{r}="","",IF(N{r}=0,"en curso",IF(N(G{r})=0,"",I{r}/G{r})))', PCT)
    calc(ws.cell(row=r, column=11), f'=IF(A{r}="","",SUM(B{r}:F{r}))', EUR)
    calc(ws.cell(row=r, column=12), f'=IF(A{r}="","",YEAR(A{r}))', "0")
    calc(ws.cell(row=r, column=13), f'=IF(A{r}="","",COUNTIFS({MVR("B")},A{r},{MVR("H")},"Por revisar"))', "0")
    calc(ws.cell(row=r, column=14), f'=IF(A{r}="","",IF(A{r}<DATE(YEAR(TODAY()),MONTH(TODAY()),1),1,0))', "0")
    for c in range(7, 15):
        ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F2F2")
ws.conditional_formatting.add(f"M6:M{LAST}", FormulaRule(formula=['N($M6)>0'], fill=PatternFill("solid", fgColor="FFF3B0")))
widths(ws, {"A": 11, "B": 13, "C": 12, "D": 12, "E": 11, "F": 12, "G": 13, "H": 13, "I": 13, "J": 10, "K": 15, "L": 7, "M": 10, "N": 10})
ws.freeze_panes = "B6"

# ================================================================ CASA
ws = ws_casa
title(ws, "CASA · gastos de la vivienda", "Gastos de la vivienda por concepto, últimos 12 meses. Sale de las categorías del grupo «Casa».", 16)
casa_cats = [c for c, t, g_ in CFG["categorias"] if g_ == "Casa"]
ws["A3"] = "Si añades una categoría del grupo «Casa» en CATEGORÍAS, añade aquí su fila."
ws["A3"].font = font(size=9, color="666666")
header(ws, 4, ["Concepto"] + [""] * 12 + ["Total 12 meses", "Media al mes", "Último mes"])
LASTMES = f"INDEX({DR('A')},{NDAT})"
for j in range(12):
    c = ws.cell(row=4, column=2 + j, value=f"=EDATE({LASTMES},{j-11})")
    c.number_format = "mmm yy"; c.font = font(size=9, bold=True, color="FFFFFF")
C0 = 5
for i, cat in enumerate(casa_cats):
    r = C0 + i
    ws.cell(row=r, column=1, value=cat).font = font()
    for j in range(12):
        col = 2 + j; L = ws.cell(row=4, column=col).column_letter
        calc(ws.cell(row=r, column=col), f"=-SUMIFS({MVR('E')},{MVR('G')},$A{r},{MVR('B')},{L}$4)", EUR)
    calc(ws.cell(row=r, column=14), f"=SUM(B{r}:M{r})", EUR0, bold=True)
    calc(ws.cell(row=r, column=15), f"=IFERROR(SUMIFS(B{r}:M{r},$B$4:$M$4,\"<\"&DATE(YEAR(TODAY()),MONTH(TODAY()),1))/$B${C0+len(casa_cats)+2},0)", EUR0)
    calc(ws.cell(row=r, column=16), f"=M{r}", EUR)
CT = C0 + len(casa_cats)
ws.cell(row=CT, column=1, value="Total casa").font = font(bold=True)
for col in range(2, 17):
    L = ws.cell(row=4, column=col).column_letter
    calc(ws.cell(row=CT, column=col), f"=SUM({L}{C0}:{L}{CT-1})", EUR0, bold=True)
    ws.cell(row=CT, column=col).fill = PatternFill("solid", fgColor=SECT)
ws.cell(row=CT, column=1).fill = PatternFill("solid", fgColor=SECT)
ws.cell(row=CT + 2, column=1, value="Meses con datos en la ventana").font = font(size=9, color="666666")
calc(ws.cell(row=CT + 2, column=2), f"=COUNTIFS($B$4:$M$4,\">=\"&'MIS DATOS'!$A$6,$B$4:$M$4,\"<\"&DATE(YEAR(TODAY()),MONTH(TODAY()),1))", "0")
ws.cell(row=CT + 3, column=1, value="· La media usa solo meses completos con datos. Agua y algunos recibos son bimestrales: mira la media, no un mes suelto.").font = font(size=9, color="555555")
CASA_MEDIA = f"CASA!$O${CT}"
bc = BarChart(); bc.type = "col"; bc.grouping = "stacked"; bc.overlap = 100; bc.title = "Gastos de casa por mes"
bc.add_data(Reference(ws, min_col=1, max_col=13, min_row=C0, max_row=CT - 1), from_rows=True, titles_from_data=True)
bc.set_categories(Reference(ws, min_col=2, max_col=13, min_row=4, max_row=4))
bc.y_axis.number_format = '#,##0'; bc.height = 8; bc.width = 22
ws.add_chart(bc, f"A{CT+5}")
widths(ws, {"A": 26, **{ws.cell(row=4, column=c).column_letter: 10 for c in range(2, 14)}, "N": 14, "O": 12, "P": 11})

# ================================================================ VIVIENDA
ws = ws_viv
title(ws, "VIVIENDA · ¿Comprar ahora o esperar?", "Compara comprar ya con seguir de alquiler unos años y comprar después, con el mismo esfuerzo al mes. Datos ficticios: pon los tuyos.", 10)
section(ws, 4, "SUPUESTOS", 4)
viv_in = [
    ("tas", "Precio de la vivienda", VIV["precio"], EUR0, "Ficticio. Lo que costaría hoy"),
    ("parte", "Parte que ya tienes", VIV["parte"], PCT, "0 % si la compras entera"),
    ("S", "Ahorros que pones en la compra", VIV["ahorros"], EUR0, "El resto lo financia la hipoteca"),
    ("ajd", "Impuestos de la compra", VIV["impuestos"], "0.00%", "ITP si es de segunda mano (cambia según la comunidad autónoma); IVA y AJD si es nueva. Confírmalo"),
    ("not", "Notaría + registro + gestoría", VIV["notaria"], EUR0, "Estimación; pide presupuesto"),
    ("ctas", "Coste de la tasación", VIV["tasacion"], EUR0, "La paga el comprador"),
    ("tipo", "Tipo de interés de la hipoteca", VIV["tipo"], "0.00%", "Ficticio: compara ofertas"),
    ("plazo", "Plazo de la hipoteca que pedirías (años)", VIV["plazo"], "0", ""),
    ("emax", "Edad máxima al terminar de pagar", VIV["edad_max"], "0", "Muchos bancos ponen el límite entre 70 y 75: pregúntalo. Si se pasa, el plazo se acorta solo"),
    ("alq", "Alquiler que pagas mientras no compras (al mes)", VIV["alquiler"], EUR0, "Si ya vives en tu casa, pon 0"),
    ("W", "Años de espera (escenario B)", VIV["espera"], "0", ""),
    ("H", "Horizonte de comparación (años)", VIV["horizonte"], "0", "Máximo 30"),
]
PI_ = {}
for i, (k, lab, val, fmt, note) in enumerate(viv_in):
    r = 5 + i
    ws.cell(row=r, column=1, value=lab).font = font()
    inp(ws.cell(row=r, column=2), val, fmt)
    ws.cell(row=r, column=3, value=note).font = font(size=9, color="666666")
    PI_[k] = f"$B${r}"
r = 5 + len(viv_in)
ws.cell(row=r, column=1, value="Revalorización anual de la vivienda").font = font()
link(ws.cell(row=r, column=2), f"={PP['rev']}", PCT1); PI_["rev"] = f"$B${r}"
ws.cell(row=r, column=3, value="Viene de PARÁMETROS").font = font(size=9, color="666666")
r += 1
ws.cell(row=r, column=1, value="Rentabilidad si el dinero sigue invertido").font = font()
link(ws.cell(row=r, column=2), f"={PP['r_inv']}", PCT1); PI_["r"] = f"$B${r}"
ws.cell(row=r, column=3, value="Viene de PARÁMETROS").font = font(size=9, color="666666")

g = PI_
g["edad"] = PP["edad"]
# Scenario A
ra = 20
section(ws, ra, "ESCENARIO A · COMPRAR YA", 4)
section_b_col = 6
A_rows = [
    ("pA", "Lo que compras (precio × lo que te falta)", f"={g['tas']}*(1-{g['parte']})", EUR0),
    ("ajdA", "Impuestos", "=B{pA}*{ajd}", EUR0),
    ("gA", "Gastos totales (impuestos + notaría + tasación)", "=B{ajdA}+{not}+{ctas}", EUR0),
    ("nA", "Dinero necesario", "=B{pA}+B{gA}", EUR0),
    ("hA", "Hipoteca", "=MAX(0,B{nA}-{S})", EUR0),
    ("cashA", "Sale de tus ahorros", "=B{nA}-B{hA}", EUR0),
    ("pctA", "% financiado de lo que compras", "=IFERROR(B{hA}/B{pA},0)", PCT1),
    ("ltvA", "% sobre el valor de la vivienda (lo que mira el banco)", "=IFERROR(B{hA}/{tas},0)", PCT1),
    ("plA", "Plazo real (con tu edad)", "=MAX(1,MIN({plazo},{emax}-{edad}))", "0"),
    ("cA", "Cuota mensual", "=IF(B{hA}=0,0,PMT({tipo}/12,B{plA}*12,-B{hA}))", EUR0),
    ("iA", "Intereses totales", "=B{cA}*B{plA}*12-B{hA}", EUR0),
]
RA = {}
for i, (k, lab, f, fmt) in enumerate(A_rows):
    RA[k] = ra + 1 + i
for i, (k, lab, f, fmt) in enumerate(A_rows):
    r = ra + 1 + i
    ws.cell(row=r, column=1, value=lab).font = font()
    calc(ws.cell(row=r, column=2), f.format(**RA, **g), fmt, bold=k in ("hA", "cA", "ltvA"))

# Scenario B (columns F:G)
B_rows = [
    ("vW", "Precio de la vivienda cuando compres", "={tas}*(1+{rev})^{W}", EUR0),
    ("pB", "Lo que compras entonces", "=G{vW}*(1-{parte})", EUR0),
    ("gB", "Gastos totales", "=G{pB}*{ajd}+{not}+{ctas}", EUR0),
    ("nB", "Dinero necesario", "=G{pB}+G{gB}", EUR0),
    ("bote", "Lo que habrás juntado (ahorros + lo que sobra de la cuota tras el alquiler)", "={S}*(1+{r})^{W}+(B{cA}-{alq})*12*IFERROR(((1+{r})^{W}-1)/{r},{W})", EUR0),
    ("hB", "Hipoteca", "=MAX(0,G{nB}-G{bote})", EUR0),
    ("cashB", "Sale de lo que has juntado", "=G{nB}-G{hB}", EUR0),
    ("ltvB", "% sobre el valor de la vivienda", "=IFERROR(G{hB}/G{vW},0)", PCT1),
    ("plB", "Plazo real (con tu edad entonces)", "=MAX(1,MIN({plazo},{emax}-{edad}-{W}))", "0"),
    ("cB", "Cuota mensual", "=IF(G{hB}=0,0,PMT({tipo}/12,G{plB}*12,-G{hB}))", EUR0),
    ("sobre", "Te cuesta de más por esperar", "=G{pB}-B{pA}", EUR0),
]
ws.cell(row=ra, column=6, value="ESCENARIO B · ESPERAR").font = font(size=11, bold=True, color=NAVY)
for c in range(6, 8): ws.cell(row=ra, column=c).fill = PatternFill("solid", fgColor=SECT)
RB = {}
for i, (k, lab, f, fmt) in enumerate(B_rows):
    RB[k] = ra + 1 + i
allk = {**RA, **RB, **g}
for i, (k, lab, f, fmt) in enumerate(B_rows):
    r = ra + 1 + i
    ws.cell(row=r, column=6, value=lab).font = font()
    calc(ws.cell(row=r, column=7), f.format(**allk), fmt, bold=k in ("hB", "cB", "sobre"))

# Year table
T0 = 46
section(ws, T0 - 2, "EVOLUCIÓN AÑO A AÑO (patrimonio ligado a la vivienda: vivienda − deuda + lo invertido)", 9)
header(ws, T0 - 1, ["Año", "Valor de la vivienda", "A · Deuda", "A · Invertido", "A · Patrimonio", "B · Deuda", "B · Invertido", "B · Patrimonio", "A − B"])
cA = f"$B${RA['cA']}"; plA = f"$B${RA['plA']}"; plB = f"$G${RB['plB']}"; hA = f"$B${RA['hA']}"; nA = f"$B${RA['nA']}"
cB = f"$G${RB['cB']}"; hB = f"$G${RB['hB']}"; nB = f"$G${RB['nB']}"; bote = f"$G${RB['bote']}"
for y in range(31):
    r = T0 + y
    ws.cell(row=r, column=1, value=y).font = font()
    calc(ws.cell(row=r, column=2), f"={g['tas']}*(1+{g['rev']})^A{r}", EUR0)
    calc(ws.cell(row=r, column=3), f"=IF(A{r}>={plA},0,MAX(0,FV({g['tipo']}/12,A{r}*12,{cA},-{hA})))", EUR0)
    if y == 0:
        calc(ws.cell(row=r, column=4), f"=MAX(0,{g['S']}-{nA})", EUR0)
        calc(ws.cell(row=r, column=7), f"=IF({g['W']}=0,MAX(0,{bote}-{nB}),{g['S']})", EUR0)
    else:
        calc(ws.cell(row=r, column=4), f"=D{r-1}*(1+{g['r']})+IF(A{r}>{plA},{cA}*12,0)", EUR0)
        calc(ws.cell(row=r, column=7),
             f"=IF(A{r}<{g['W']},G{r-1}*(1+{g['r']})+({cA}-{g['alq']})*12,IF(A{r}={g['W']},MAX(0,{bote}-{nB}),"
             f"G{r-1}*(1+{g['r']})+IF(A{r}-{g['W']}<={plB},({cA}-{cB})*12,{cA}*12)))", EUR0)
    calc(ws.cell(row=r, column=5), f"=B{r}-C{r}+D{r}", EUR0)
    calc(ws.cell(row=r, column=6), f"=IF(A{r}<{g['W']},0,IF(A{r}-{g['W']}>={plB},0,MAX(0,FV({g['tipo']}/12,(A{r}-{g['W']})*12,{cB},-{hB}))))", EUR0)
    calc(ws.cell(row=r, column=8), f"=IF(A{r}<{g['W']},B{r}*{g['parte']}+G{r},B{r}-F{r}+G{r})", EUR0)
    calc(ws.cell(row=r, column=9), f"=E{r}-H{r}", EUR)
TE = T0 + 30

# Result
rr = 33
section(ws, rr, "RESULTADO", 9)
ws.cell(row=rr + 1, column=1, value="Años de comparación").font = font()
link(ws.cell(row=rr + 1, column=2), f"={g['H']}", "0")
ws.cell(row=rr + 2, column=1, value="Patrimonio si compras ya (A)").font = font()
calc(ws.cell(row=rr + 2, column=2), f"=INDEX($E${T0}:$E${TE},MIN(30,{g['H']})+1)", EUR0, bold=True)
ws.cell(row=rr + 3, column=1, value="Patrimonio si esperas (B)").font = font()
calc(ws.cell(row=rr + 3, column=2), f"=INDEX($H${T0}:$H${TE},MIN(30,{g['H']})+1)", EUR0, bold=True)
ws.cell(row=rr + 4, column=1, value="Diferencia (A − B)").font = font()
calc(ws.cell(row=rr + 4, column=2), f"=B{rr+2}-B{rr+3}", EUR, bold=True)
calc(ws.cell(row=rr + 5, column=1),
     f'=IF(B{rr+4}>=0,"✅ Comprar ya sale mejor por "&TEXT(B{rr+4},"#,##0")&" € a "&B{rr+1}&" años",'
     f'"⏳ Esperar sale mejor por "&TEXT(-B{rr+4},"#,##0")&" € a "&B{rr+1}&" años")', bold=True, size=12, color=NAVY)
notes = [
    "Mismo esfuerzo al mes en los dos: en B pagas el alquiler e inviertes lo que te sobra de la cuota de A.",
    "Lo que decide casi siempre: revalorización de la vivienda e intereses de la hipoteca frente a lo que rinde el dinero invertido y lo que se va en alquiler.",
    "Riesgos de esperar que no salen en la cuenta: que los precios o los tipos suban más de lo previsto.",
    "No es asesoramiento fiscal ni financiero: confirma impuestos y gastos con una gestoría antes de firmar.",
]
calc(ws.cell(row=rr + 6, column=1), f'=IF(OR($B${RA["plA"]}<{g["plazo"]},$G${RB["plB"]}<{g["plazo"]}),"⚠️ Por tu edad el plazo se acorta (A: "&$B${RA["plA"]}&" años, B: "&$G${RB["plB"]}&" años) y la cuota sube","")', bold=True, color="C62828")
for i, n in enumerate(notes):
    ws.cell(row=rr + 7 + i, column=1, value="· " + n).font = font(size=9, color="555555")
widths(ws, {"A": 46, "B": 15, "C": 14, "D": 14, "E": 15, "F": 44, "G": 15, "H": 15, "I": 13})
COMPRA_CASH_A = f"{PI}$B${RA['cashA']}"; COMPRA_CASH_B = f"{PI}$G${RB['cashB']}"
COMPRA_CA = f"{PI}$B${RA['cA']}"; COMPRA_CB = f"{PI}$G${RB['cB']}"
COMPRA_HA = f"{PI}$B${RA['hA']}"; COMPRA_HB = f"{PI}$G${RB['hB']}"
COMPRA_PLA = f"{PI}$B${RA['plA']}"; COMPRA_PLB = f"{PI}$G${RB['plB']}"

ch = LineChart(); ch.title = "Patrimonio ligado a la vivienda: comprar ya vs esperar"
ch.y_axis.title = "€"; ch.x_axis.title = "Años"; ch.height = 8; ch.width = 17
ch.add_data(Reference(ws, min_col=5, min_row=T0 - 1, max_row=T0 + 20), titles_from_data=True)
ch.add_data(Reference(ws, min_col=8, min_row=T0 - 1, max_row=T0 + 20), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=T0, max_row=T0 + 20))
ch.y_axis.number_format = '#,##0'
ws.add_chart(ch, "K4")

# ================================================================ INMUEBLES
ws = ws_inm
title(ws, "INMUEBLES", "Lo que tienes en ladrillo. Suma a tu patrimonio; a la libertad financiera solo si da renta. Valores estimados: cámbialos.", 8)
header(ws, 4, ["Inmueble", "Valor estimado (100 %)", "Tu %", "Tu parte", "Renta anual", "Gastos anuales", "Revalorización anual", "Nota"])
inm = [("Vivienda (la de VIVIENDA)", f"={PI}{g['tas']}", f"={PI}{g['parte']}", 0, 0, f"={PP['rev']}", "Valor y % vienen de VIVIENDA: 0 % mientras no la compres", True)] + \
      [(None, None, None, None, None, None, None, False)] * 4
for i, (n, v, pc, ren, gas, rv, note, linked) in enumerate(inm):
    r = 5 + i
    inp(ws.cell(row=r, column=1), n)
    if linked:
        link(ws.cell(row=r, column=2), v, EUR0); link(ws.cell(row=r, column=3), pc, PCT)
    else:
        inp(ws.cell(row=r, column=2), v, EUR0); inp(ws.cell(row=r, column=3), pc, PCT)
    inp(ws.cell(row=r, column=5), ren, EUR0); inp(ws.cell(row=r, column=6), gas, EUR0)
    if isinstance(rv, str): link(ws.cell(row=r, column=7), rv, PCT1)
    else: inp(ws.cell(row=r, column=7), rv, PCT1)
    calc(ws.cell(row=r, column=4), f'=IF(B{r}="",0,B{r}*C{r})', EUR0)
    ws.cell(row=r, column=8, value=note).font = font(size=9, color="666666")
ws["A10"] = "Total"; ws["A10"].font = font(bold=True)
calc(ws["D10"], "=SUM(D5:D9)", EUR0, bold=True)
calc(ws["E10"], "=SUM(E5:E9)", EUR0, bold=True)
calc(ws["F10"], "=SUM(F5:F9)", EUR0, bold=True)
ws["A11"] = "Renta neta anual"; ws["A11"].font = font()
calc(ws["D11"], "=E10-F10", EUR, bold=True)
ws["A13"] = "· La renta neta, si algún día la hay, se suma a la pensión en la fase de jubilación."
ws["A14"] = "· Cada inmueble lleva su revalorización: un local o un garaje no se mueven como la vivienda."
for c in ("A13", "A14"): ws[c].font = font(size=9, color="555555")
widths(ws, {"A": 22, "B": 20, "C": 8, "D": 14, "E": 12, "F": 13, "G": 14, "H": 52})
INM_TOT = f"{IN}$D$10"; INM_REN = f"{IN}$D$11"

# ================================================================ VIAJES
ws = ws_via
title(ws, "VIAJES · tu cuenta de viajes", "Aportas cada mes a una cuenta de viajes y de ahí pagas tu parte. Cada viaje te dice cuánto queda en la cuenta y si llega.", 12)
section(ws, 3, "TU CUENTA DE VIAJES", 3)
ws["A4"] = "Aportación mensual"; inp(ws["B4"], VJ["aportacion"], EUR0)
ws["A5"] = "Primera aportación (mes)"; inp(ws["B5"], d_(VJ["primera_aportacion"]), "mmm yyyy")
ws["A6"] = "Saldo que ya había en la cuenta"; inp(ws["B6"], VJ["saldo_previo"], EUR0)
ws["A7"] = "Año que miras"; inp(ws["B7"], ANIO, "0")
ws["A8"] = "Aportado ese año"; ws["A9"] = "Pagado de la cuenta ese año (tu parte)"; ws["A10"] = "Saldo a final de año"
ws["A11"] = "Viajes de ese año"; ws["A12"] = "Aviso"
for r in range(4, 13): ws.cell(row=r, column=1).font = font()
ws["C4"] = "€ al mes"
ws["C4"].font = font(size=9, color="666666")
VT0, VT1 = 16, 65
TR = lambda col: f"${col}${VT0}:${col}${VT1}"
MESES_HASTA = lambda fecha: f"MAX(0,(YEAR({fecha})-YEAR($B$5))*12+MONTH({fecha})-MONTH($B$5)+1)"
calc(ws["B8"], f"=$B$4*({MESES_HASTA('DATE($B$7,12,1)')}-{MESES_HASTA('DATE($B$7-1,12,1)')})", EUR0)
calc(ws["B9"], f'=SUMIFS({TR("I")},{TR("F")},"Cuenta de viajes",{TR("B")},">="&DATE($B$7,1,1),{TR("B")},"<="&DATE($B$7,12,31))', EUR0)
calc(ws["B10"], f'=$B$6+$B$4*{MESES_HASTA("DATE($B$7,12,1)")}-SUMIFS({TR("I")},{TR("F")},"Cuenta de viajes",{TR("B")},"<="&DATE($B$7,12,31))', EUR0, bold=True, size=12)
calc(ws["B11"], f'=COUNTIFS({TR("B")},">="&DATE($B$7,1,1),{TR("B")},"<="&DATE($B$7,12,31))', "0")
calc(ws["B12"], f'=IFERROR(IF(MIN({TR("J")})<0,"⚠️ No llega para "&INDEX({TR("A")},MATCH(MIN({TR("J")}),{TR("J")},0))&": faltan "&TEXT(-MIN({TR("J")}),"#,##0")&" €","✅ Llega para todos los viajes apuntados"),"")', bold=True, color="8A6D00")
header(ws, VT0 - 1, ["Viaje", "Fecha", "Coste total", "Personas", "Tu parte", "Se paga con", "Estado", "Coste real (tu parte)", "Cuenta", "Saldo de la cuenta tras el viaje", "¿Llega?", "Nota"])
ws.row_dimensions[VT0 - 1].height = 30
trips = [(n_, d_(f_), c_, p_, con_, est_, nota_) for n_, f_, c_, p_, con_, est_, nota_ in VJ["lista"]]
dv1 = DataValidation(type="list", formula1='"Idea,Planificando,Organizado,Hecho"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"Cuenta de viajes,Mi bolsillo"', allow_blank=True)
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
for i in range(VT1 - VT0 + 1):
    r = VT0 + i
    t = trips[i] if i < len(trips) else (None, None, None, None, None, None, None)
    inp(ws.cell(row=r, column=1), t[0]); inp(ws.cell(row=r, column=2), t[1], "dd/mm/yyyy")
    inp(ws.cell(row=r, column=3), t[2], EUR0); inp(ws.cell(row=r, column=4), t[3], "0")
    inp(ws.cell(row=r, column=6), t[4]); inp(ws.cell(row=r, column=7), t[5]); inp(ws.cell(row=r, column=8), None, EUR0)
    ws.cell(row=r, column=12, value=t[6]).font = font(size=9, color="666666")
    dv1.add(f"G{r}"); dv2.add(f"F{r}")
    calc(ws.cell(row=r, column=5), f'=IF(OR(A{r}="",C{r}=""),"",C{r}/MAX(1,N(D{r})))', EUR0)
    calc(ws.cell(row=r, column=9), f'=IF(A{r}="","",IF(H{r}<>"",H{r},E{r}))', EUR0)
    # helpers: meses de aportación hasta el viaje (M) y pagado de la cuenta hasta el viaje (N)
    calc(ws.cell(row=r, column=13), f'=IF(OR(A{r}="",B{r}=""),"",{MESES_HASTA(f"B{r}")})', "0")
    calc(ws.cell(row=r, column=14), f'=IF(OR(A{r}="",B{r}="",F{r}<>"Cuenta de viajes"),"",SUMIFS({TR("I")},{TR("F")},"Cuenta de viajes",{TR("B")},"<="&B{r}))', EUR0)
    calc(ws.cell(row=r, column=10), f'=IF(N{r}="","",$B$6+$B$4*M{r}-N{r})', EUR)
    calc(ws.cell(row=r, column=11), f'=IF(A{r}="","",IF(F{r}<>"Cuenta de viajes","— lo pagas tú",IF(OR(B{r}="",E{r}=""),"Falta fecha o coste",IF(J{r}<0,"🔴 Faltan "&TEXT(-J{r},"#,##0")&" €",IF(J{r}<$B$4,"🟡 Justo","🟢 Llega")))))')
    for c in (13, 14): ws.cell(row=r, column=c).font = font(size=8, color="AAAAAA")
ws.cell(row=VT0 - 1, column=13, value="meses aportados").font = font(size=8, color="AAAAAA")
ws.cell(row=VT0 - 1, column=14, value="pagado de la cuenta hasta aquí").font = font(size=8, color="AAAAAA")
ws.conditional_formatting.add(f"K{VT0}:K{VT1}", FormulaRule(formula=[f'LEFT($K{VT0},1)="🔴"'], fill=PatternFill("solid", fgColor="FDE2E1")))

# ---- ¿Y si aporto otra cantidad?  (etiquetas E:G combinadas, valores H..K, ayudas P..S)
LABC, VALC, HCOL = 5, [8, 9, 10, 11], [16, 17, 18, 19]
CL = lambda c: ws.cell(row=1, column=c).column_letter
ws.cell(row=3, column=LABC, value="¿Y SI APORTO OTRA CANTIDAD?").font = font(size=11, bold=True, color=NAVY)
for c in range(LABC, VALC[-1] + 1): ws.cell(row=3, column=c).fill = PatternFill("solid", fgColor=SECT)
for c, h in zip([LABC] + VALC, ["", "Lo de ahora", "Opción A", "Opción B", "Opción C"]):
    cell = ws.cell(row=4, column=c, value=h); cell.font = font(size=9, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY); cell.alignment = Alignment(horizontal="center")
for c in (6, 7): ws.cell(row=4, column=c).fill = PatternFill("solid", fgColor=NAVY)
rows_wi = ["Aportación mensual", f"Saldo de la cuenta a fin de {ANIO}", f"Saldo de la cuenta a fin de {ANIO + 1}",
           "Saldo más bajo justo después de un viaje", "¿Llegas a todos los viajes?", "Ahorro que te queda al mes",
           "% de tus ingresos que ahorras", f"Patrimonio líquido a los {JUB}", f"Libertad financiera a los {JUB}"]
for i, lab in enumerate(rows_wi):
    r = 5 + i
    ws.merge_cells(start_row=r, start_column=LABC, end_row=r, end_column=7)
    ws.cell(row=r, column=LABC, value=lab).font = font()
calc(ws.cell(row=5, column=VALC[0]), "=$B$4", EUR0, bold=True)
for j, v in enumerate(VJ["opciones"]):
    inp(ws.cell(row=5, column=VALC[1 + j]), v, EUR0)
OBS = f"IFERROR(SUMIFS({MVR('E')},{MVR('G')},\"Aportación a la cuenta de viajes\",{MVR('B')},\">=\"&'MIS DATOS'!$A$6,{MVR('B')},\"<\"&DATE(YEAR(TODAY()),MONTH(TODAY()),1))/-COUNTIFS({DR('N')},1),0)"
NJ = "'PROYECCIÓN'!$B$11"
P67 = f"INDEX('PROYECCIÓN'!$G$15:$G$45,{NJ}+1)"
FACT = f"IFERROR(((1+{PP['r_ef']})^{NJ}-1)/{PP['r_ef']},{NJ})"
INGM = f"IFERROR(AVERAGEIFS({DR('G')},{DR('N')},1),0)"
for j in range(4):
    col = VALC[j]; X = CL(col); a = f"{X}$5"; H = CL(HCOL[j])
    calc(ws.cell(row=6, column=col), f'=$B$6+{a}*{MESES_HASTA(f"DATE({ANIO},12,1)")}-SUMIFS({TR("I")},{TR("F")},"Cuenta de viajes",{TR("B")},"<="&DATE({ANIO},12,31))', EUR)
    calc(ws.cell(row=7, column=col), f'=$B$6+{a}*{MESES_HASTA(f"DATE({ANIO + 1},12,1)")}-SUMIFS({TR("I")},{TR("F")},"Cuenta de viajes",{TR("B")},"<="&DATE({ANIO + 1},12,31))', EUR)
    ws.cell(row=VT0 - 1, column=HCOL[j], value=f"saldo · {['ahora','A','B','C'][j]}").font = font(size=8, color="AAAAAA")
    for r in range(VT0, VT1 + 1):
        calc(ws.cell(row=r, column=HCOL[j]), f'=IF(N{r}="","",$B$6+{a}*M{r}-N{r})', EUR0)
        ws.cell(row=r, column=HCOL[j]).font = font(size=8, color="AAAAAA")
    calc(ws.cell(row=8, column=col), f'=IF(COUNT({H}{VT0}:{H}{VT1})=0,"",MIN({H}{VT0}:{H}{VT1}))', EUR)
    calc(ws.cell(row=9, column=col), f'=IF({X}8="","Sin viajes de la cuenta",IF({X}8<0,"🔴 No","🟢 Sí"))')
    calc(ws.cell(row=10, column=col), f"={PP['ap_ef']}+{OBS}-{a}", EUR)
    calc(ws.cell(row=11, column=col), f'=IFERROR({X}10/{INGM},"")', PCT)
    calc(ws.cell(row=12, column=col), f"={P67}+({OBS}-{a})*12*{FACT}", EUR0)
    calc(ws.cell(row=13, column=col), f"=IFERROR({X}12*{PP['ret']}/(({PP['gastos']}+{a}-{OBS})*12*(1+{PP['infl']})^{NJ}),\"\")", PCT)
    for r in range(5, 14): ws.cell(row=r, column=col).alignment = Alignment(horizontal="right")
ws.cell(row=14, column=LABC, value=f"Ahorro sobre tu media real: suma lo que ya aportas a la cuenta de viajes y resta la opción. Patrimonio y libertad a los {JUB} suponen aportar esa cantidad cada año.").font = font(size=8, color="555555")
widths(ws, {"A": 40, "B": 12, "C": 12, "D": 9, "E": 11, "F": 16, "G": 13, "H": 13, "I": 12, "J": 15, "K": 16, "L": 34, "M": 8, "N": 10, "O": 2,
            "P": 9, "Q": 9, "R": 9, "S": 9})
ws.freeze_panes = f"B{VT0}"

# ================================================================ DECISIONES
ws = ws_dec
title(ws, "DECISIONES · ¿qué pasa si…?", "Cambia las palancas y mira en MI FUTURO cómo se mueve tu patrimonio frente al plan base.", 4)
section(ws, 4, "PALANCAS", 3)
dec_in = [("lump", "Pasar ya de tu efectivo a inversión", 0, EUR0, "Si tu efectivo casi no da intereses, prueba qué pasa si inviertes una parte"),
          ("extra", "Aportación mensual a inversión", 0, EUR0, ""),
          ("rec", "Recorte de gastos mensuales", 0, EUR0, ""),
          ("invrec", "¿Inviertes lo que recortas?", "Sí", None, ""),
          ("r_esc", "Rentabilidad de la inversión en este escenario", 0.05, PCT1, ""),
          ("compra", "¿Incluir la compra de la vivienda?", "Sí", None, ""),
          ("esc", "Escenario de compra", "Comprar ya", None, "Los números salen de VIVIENDA")]
DD = {}
dvs = DataValidation(type="list", formula1='"Sí,No"'); dve = DataValidation(type="list", formula1='"Comprar ya,Esperar"')
ws.add_data_validation(dvs); ws.add_data_validation(dve)
for i, (k, lab, v, fmt, note) in enumerate(dec_in):
    r = 5 + i
    ws.cell(row=r, column=1, value=lab).font = font()
    if isinstance(v, str) and v.startswith("="): link(ws.cell(row=r, column=2), v, fmt)
    else: inp(ws.cell(row=r, column=2), v, fmt)
    ws.cell(row=r, column=3, value=note).font = font(size=9, color="666666")
    DD[k] = f"{DE}$B${r}"
    if k in ("invrec", "compra"): dvs.add(f"B{r}")
    if k == "esc": dve.add(f"B{r}")
R1 = 5 + len(dec_in) + 1
section(ws, R1, "LO QUE ENTRA Y SALE (calculado)", 3)
res = [("tc", "Dentro de cuántos años compras", f'=IF({DD["esc"]}="Comprar ya",0,{PI}{g["W"]})', "0"),
       ("cash", "Sale de tus ahorros para la compra", f'=IF({DD["compra"]}="Sí",IF({DD["esc"]}="Comprar ya",{COMPRA_CASH_A},{COMPRA_CASH_B}),0)', EUR0),
       ("hip", "Hipoteca", f'=IF({DD["compra"]}="Sí",IF({DD["esc"]}="Comprar ya",{COMPRA_HA},{COMPRA_HB}),0)', EUR0),
       ("cuota", "Cuota mensual", f'=IF({DD["compra"]}="Sí",IF({DD["esc"]}="Comprar ya",{COMPRA_CA},{COMPRA_CB}),0)', EUR0),
       ("plazo", "Años de hipoteca", f'=IF({DD["esc"]}="Comprar ya",{COMPRA_PLA},{COMPRA_PLB})', "0"),
       ("alq", "Alquiler que dejas de pagar al comprar (al mes, € de hoy)", f'=IF({DD["compra"]}="Sí",{PI}{g["alq"]},0)', EUR0)]
for i, (k, lab, f, fmt) in enumerate(res):
    r = R1 + 1 + i
    ws.cell(row=r, column=1, value=lab).font = font()
    calc(ws.cell(row=r, column=2), f, fmt)
    DD[k] = f"{DE}$B${r}"
R2 = R1 + len(res) + 2
calc(ws.cell(row=R2, column=1), f'=IF(AND({DD["compra"]}="Sí",{DD["tc"]}=0,{DD["cash"]}>{PR}$B$5+{PR}$B$6-{PP["gastos"]}*{PP["colchon"]}),"⚠️ Con tus saldos de hoy, la compra se come parte del colchón de emergencia","")', bold=True, color="C62828")
calc(ws.cell(row=R2 + 1, column=1), f'=IF(AND({DD["compra"]}="Sí",{DD["tc"]}+{DD["plazo"]}>{PP["jub"]}-{PP["edad"]}),"ℹ️ La hipoteca sigue "&({DD["tc"]}+{DD["plazo"]}-({PP["jub"]}-{PP["edad"]}))&" años después de jubilarte: la cuota se suma a lo que sacas","")', color="1565C0")
notes_d = ["· El escenario saca la compra de tu inversión y paga la cuota con lo que ibas a aportar. Desde que compras, el alquiler que ya no pagas va a tu inversión y deja de contar como gasto.",
           "· Los inmuebles no cuentan como libertad financiera: no dan para vivir salvo que den renta o los vendas."]
for i, n in enumerate(notes_d):
    ws.cell(row=R2 + 3 + i, column=1, value=n).font = font(size=9, color="555555")
widths(ws, {"A": 50, "B": 16, "C": 60})

# ================================================================ PROYECCIÓN
ws = ws_pro
title(ws, "PROYECCIÓN", "Cálculo año a año. No hace falta tocar nada aquí. Si «Inversión con decisiones» sale negativa, ese dinero sale del efectivo.", 22)
section(ws, 3, "PUNTO DE PARTIDA (último mes de MIS DATOS)", 3)
start = [("mes", "Último mes registrado", f"=INDEX({DR('A')},{NDAT})", "mmm yyyy"),
         ("ef0", f"Efectivo ({' + '.join(CU)})", f"=INDEX({DR('B')},{NDAT})+INDEX({DR('C')},{NDAT})+INDEX({DR('D')},{NDAT})", EUR0),
         ("inv0", "Inversión", f"=INDEX({DR('E')},{NDAT})", EUR0),
         ("pen0", "Plan de pensiones", f"=INDEX({DR('F')},{NDAT})", EUR0),
         ("liq0", "Patrimonio líquido", "=B5+B6+B7", EUR0),
         ("inm0", "Tus inmuebles (tu parte)", f"={INM_TOT}", EUR0),
         ("tot0", "Patrimonio total con inmuebles", "=B8+B9", EUR0),
         ("tj", "Años hasta la jubilación", f"={PP['jub']}-{PP['edad']}", "0")]
S0 = {}
for i, (k, lab, f, fmt) in enumerate(start):
    r = 4 + i
    ws.cell(row=r, column=1, value=lab).font = font()
    link(ws.cell(row=r, column=2), f, fmt)
    S0[k] = f"{PR}$B${r}"
TJ = "$B$11"
H0 = 14; R0 = 15; RN = R0 + 30
section(ws, 13, "FASE DE AHORRO (hasta la jubilación)", 21)
header(ws, H0, ["t", "Edad", "Año", "Efectivo", "Inversión", "Pensiones", "Total líquido", "Total en € de hoy",
                "Lo que tenías hoy", "Lo que ahorras", "Lo que generan tus inversiones", "Gastos anuales",
                "% libertad financiera", "Inversión con decisiones", "Total con decisiones", "Gastos con decisiones", "% libertad con decisiones",
                "Inmuebles (sin decisiones)", "Inmuebles netos de hipoteca (con decisiones)", "Total con inmuebles (sin decisiones)", "Total con inmuebles (con decisiones)", "Efectivo con decisiones"])
ws.row_dimensions[H0].height = 42
pr = PP; dd = DD
def W_(r, f):  # blank beyond retirement
    return f"=IF($A{r}>{TJ},\"\",{f})"
for t in range(31):
    r = R0 + t
    ws.cell(row=r, column=1, value=t).font = font(color="888888")
    calc(ws.cell(row=r, column=2), W_(r, f"{pr['edad']}+A{r}"), "0")
    calc(ws.cell(row=r, column=3), W_(r, f"{pr['anio']}+A{r}"), "0")
    grow = f"*12*(1+{pr['g']})^(A{r}-1)"
    if t == 0:
        calc(ws.cell(row=r, column=4), "=$B$5", EUR0)
        calc(ws.cell(row=r, column=5), "=$B$6", EUR0)
        calc(ws.cell(row=r, column=6), "=$B$7", EUR0)
        calc(ws.cell(row=r, column=10), "=0", EUR0)
        calc(ws.cell(row=r, column=14), f"=E{r}+{dd['lump']}-IF({dd['tc']}=0,{dd['cash']},0)", EUR0)
    else:
        calc(ws.cell(row=r, column=4), W_(r, f"D{r-1}*(1+{pr['r_ef']})+{pr['ap_ef']}{grow}"), EUR0)
        calc(ws.cell(row=r, column=5), W_(r, f"E{r-1}*(1+{pr['r_inv']})+{pr['ap_inv']}{grow}"), EUR0)
        calc(ws.cell(row=r, column=6), W_(r, f"F{r-1}*(1+{pr['r_pen']})+{pr['ap_pen']}{grow}"), EUR0)
        calc(ws.cell(row=r, column=10), W_(r, f"J{r-1}+({pr['ap_ef']}+{pr['ap_inv']}+{pr['ap_pen']}){grow}"), EUR0)
        calc(ws.cell(row=r, column=14), W_(r,
             f'N{r-1}*(1+{dd["r_esc"]})+({pr["ap_inv"]}+{dd["extra"]}+IF({dd["invrec"]}="Sí",{dd["rec"]},0)){grow}'
             f'-IF(A{r}={dd["tc"]},{dd["cash"]},0)-IF(AND(A{r}>{dd["tc"]},A{r}<={dd["tc"]}+{dd["plazo"]}),{dd["cuota"]}*12,0)'
             f'+IF(A{r}>{dd["tc"]},{dd["alq"]}*12*(1+{pr["infl"]})^A{r},0)'), EUR0)
    calc(ws.cell(row=r, column=7), W_(r, f"D{r}+E{r}+F{r}"), EUR0)
    calc(ws.cell(row=r, column=8), W_(r, f"G{r}/(1+{pr['infl']})^A{r}"), EUR0)
    calc(ws.cell(row=r, column=9), W_(r, f"$G${R0}"), EUR0)
    calc(ws.cell(row=r, column=11), W_(r, f"G{r}-I{r}-J{r}"), EUR0)
    calc(ws.cell(row=r, column=12), W_(r, f"{pr['gastos']}*12*(1+{pr['infl']})^A{r}"), EUR0)
    calc(ws.cell(row=r, column=13), W_(r, f"G{r}*{pr['ret']}/L{r}"), PCT)
    calc(ws.cell(row=r, column=22), W_(r, f"D{r}-{dd['lump']}*(1+{pr['r_ef']})^A{r}"), EUR0)
    calc(ws.cell(row=r, column=15), W_(r, f"V{r}+N{r}+F{r}"), EUR0)
    calc(ws.cell(row=r, column=16), W_(r, f"({pr['gastos']}-{dd['rec']})*12*(1+{pr['infl']})^A{r}-IF(A{r}>{dd['tc']},{dd['alq']}*12*(1+{pr['infl']})^A{r},0)"), EUR0)
    calc(ws.cell(row=r, column=17), W_(r, f"O{r}*{pr['ret']}/P{r}"), PCT)
    tas_t = f"{PI}{g['tas']}*(1+{pr['rev']})^A{r}"
    calc(ws.cell(row=r, column=18), W_(r, f"SUMPRODUCT({IN}$D$5:$D$9,(1+{IN}$G$5:$G$9)^A{r})"), EUR0)
    calc(ws.cell(row=r, column=19), W_(r,
         f'R{r}-{tas_t}*{PI}{g["parte"]}'
         f'+IF(AND({dd["compra"]}="Sí",A{r}>={dd["tc"]}),{tas_t}-IF(A{r}-{dd["tc"]}>={dd["plazo"]},0,MAX(0,FV({PI}{g["tipo"]}/12,(A{r}-{dd["tc"]})*12,{dd["cuota"]},-{dd["hip"]}))),{tas_t}*{PI}{g["parte"]})'), EUR0)
    calc(ws.cell(row=r, column=20), W_(r, f"G{r}+R{r}"), EUR0)
    calc(ws.cell(row=r, column=21), W_(r, f"O{r}+S{r}"), EUR0)

# ---- fase de jubilación
J0 = RN + 3   # section row
JH = J0 + 1; J1 = JH + 1; JN = J1 + 40
section(ws, J0, "FASE DE JUBILACIÓN (sacas cada año lo que no cubren la pensión y las rentas)", 21)
header(ws, JH, ["k", "Edad", "Año", "Capital al empezar", "Gastos del año", "Pensión del año", "Rentas inmuebles",
                "Lo que sacas", "Capital al acabar", "Capital al empezar (con decisiones)", "Gastos (con decisiones)",
                "Cuota hipoteca pendiente", "Lo que sacas (con decisiones)", "Capital al acabar (con decisiones)"])
ws.row_dimensions[JH].height = 42
for k in range(41):
    r = J1 + k
    tt = f"({TJ}+A{r})"
    ws.cell(row=r, column=1, value=k).font = font(color="888888")
    calc(ws.cell(row=r, column=2), f"={pr['jub']}+A{r}", "0")
    calc(ws.cell(row=r, column=3), f"={pr['anio']}+{tt}", "0")
    if k == 0:
        calc(ws.cell(row=r, column=4), f"=INDEX($G${R0}:$G${RN},{TJ}+1)", EUR0)
        calc(ws.cell(row=r, column=10), f"=INDEX($O${R0}:$O${RN},{TJ}+1)", EUR0)
    else:
        calc(ws.cell(row=r, column=4), f"=I{r-1}", EUR0)
        calc(ws.cell(row=r, column=10), f"=N{r-1}", EUR0)
    calc(ws.cell(row=r, column=5), f"={pr['gastos']}*12*(1+{pr['infl']})^{tt}", EUR0)
    calc(ws.cell(row=r, column=6), f"={pr['pension']}*14*(1+{pr['infl']})^{tt}", EUR0)
    calc(ws.cell(row=r, column=7), f"={INM_REN}*(1+{pr['infl']})^{tt}", EUR0)
    calc(ws.cell(row=r, column=8), f"=MAX(0,E{r}-F{r}-G{r})", EUR0)
    calc(ws.cell(row=r, column=9), f"=MAX(0,D{r}*(1+{pr['r_jub']})-H{r})", EUR0)
    calc(ws.cell(row=r, column=11), f"=({pr['gastos']}-{dd['rec']})*12*(1+{pr['infl']})^{tt}-IF({tt}>{dd['tc']},{dd['alq']}*12*(1+{pr['infl']})^{tt},0)", EUR0)
    calc(ws.cell(row=r, column=12), f"=IF(AND({tt}>={dd['tc']},{tt}<{dd['tc']}+{dd['plazo']}),{dd['cuota']}*12,0)", EUR0)
    calc(ws.cell(row=r, column=13), f"=MAX(0,K{r}+L{r}-F{r}-G{r})", EUR0)
    calc(ws.cell(row=r, column=14), f"=MAX(0,J{r}*(1+{pr['r_jub']})-M{r})", EUR0)
ws.cell(row=JN + 2, column=1, value="Te dura hasta los (sin decisiones)").font = font(bold=True)
calc(ws.cell(row=JN + 2, column=2), f"={pr['jub']}+COUNTIF($I${J1}:$I${JN},\">0\")", "0", bold=True)
ws.cell(row=JN + 3, column=1, value="Te dura hasta los (con decisiones)").font = font(bold=True)
calc(ws.cell(row=JN + 3, column=2), f"={pr['jub']}+COUNTIF($N${J1}:$N${JN},\">0\")", "0", bold=True)
DURA_B = f"{PR}$B${JN+2}"; DURA_C = f"{PR}$B${JN+3}"; DURA_MAX = 40
widths(ws, {"A": 34, "B": 14, "C": 7, "D": 12, "E": 12, "F": 12, "G": 13, "H": 13, "I": 12, "J": 14,
            "K": 14, "L": 12, "M": 13, "N": 14, "O": 13, "P": 12, "Q": 11, "R": 13, "S": 15, "T": 14, "U": 14, "V": 13})
ws.freeze_panes = f"D{R0}"
col = lambda c: f"{PR}${c}${R0}:${c}${RN}"
atj = lambda c: f"INDEX({col(c)},{PR}{TJ}+1)"

# ================================================================ MI FUTURO
ws = ws_fut
title(ws, "MI FUTURO FINANCIERO", "", 9)
calc(ws["A2"], f'="Datos ficticios de ejemplo a "&TEXT({S0["mes"]},"mmmm yyyy")&" · estimados: precio de la vivienda y pensión"', size=9, color="555555")
calc(ws["A8"], f'="🏠 Casa: "&TEXT({CASA_MEDIA},"#,##0")&" €/mes de media   ·   🧾 Por revisar: "&COUNTIFS({MVR("H")},"Por revisar",{MVR("B")},">="&\'MIS DATOS\'!$A$6)&" movimientos desde "&TEXT(\'MIS DATOS\'!$A$6,"mmmm yyyy")&" (cuentan cuando les pongas categoría)"', size=9, bold=True, color="8A6D00")
LIB = f"{PR}$M${R0}"
kpi = [("A", "TU NIVEL"), ("C", "PATRIMONIO LÍQUIDO HOY"), ("E", "LIBERTAD FINANCIERA HOY")]
for c, l in kpi:
    ws[f"{c}4"] = l; ws[f"{c}4"].font = font(size=8, bold=True, color="666666")
calc(ws["H4"], f'="CUENTA DE VIAJES · FIN DE "&{V}$B$7', size=8, bold=True, color="666666")
calc(ws["A5"], f"=INDEX({NIV},MATCH({LIB},{UMB},1))", bold=True, size=16, color=NAVY)
calc(ws["C5"], f"={S0['liq0']}", EUR0, bold=True, size=16)
calc(ws["E5"], f"={LIB}", PCT, bold=True, size=16, color=GOOD)
calc(ws["H5"], f"={V}$B$10", EUR0, bold=True, size=16)
nxt_u = f"INDEX({UMB},MATCH({LIB},{UMB},1)+1)"
nxt_n = f"INDEX({NIV},MATCH({LIB},{UMB},1)+1)"
cur_u = f"INDEX({UMB},MATCH({LIB},{UMB},1))"
calc(ws["A6"], f'=IFERROR(REPT("█",ROUND(({LIB}-{cur_u})/({nxt_u}-{cur_u})*10,0))&REPT("░",10-ROUND(({LIB}-{cur_u})/({nxt_u}-{cur_u})*10,0)),"██████████")', color=GOOD)
calc(ws["A7"], f'=IFERROR("Siguiente: "&{nxt_n}&" al "&TEXT({nxt_u},"0%"),"Nivel máximo")', size=8, color="666666")
calc(ws["C6"], f'="Con tus inmuebles: "&TEXT({S0["tot0"]},"#,##0")&" €"', size=9, color="555555")
calc(ws["C7"], f'="{CU[0]} "&TEXT(INDEX({DR("B")},{NDAT}),"#,##0")&" · {CU[1]} "&TEXT(INDEX({DR("C")},{NDAT}),"#,##0")&" · {CU[2]} "&TEXT(INDEX({DR("D")},{NDAT}),"#,##0")', size=8, color="666666")
calc(ws["E6"], f'=REPT("█",ROUND(MIN(1,{LIB})*20,0))&REPT("░",20-ROUND(MIN(1,{LIB})*20,0))', color=GOOD)
calc(ws["E7"], f'="Tu dinero cubre el "&TEXT({LIB},"0%")&" de tus gastos (regla del "&TEXT({PP["ret"]},"0%")&")"', size=8, color="666666")
calc(ws["H6"], f'="Aportas "&TEXT({V}$B$4,"#,##0")&" €/mes · "&{V}$B$11&" viajes en "&{V}$B$7', size=9, color="555555")
calc(ws["H7"], f'={V}$B$12', size=8, color="8A6D00")
ws.row_dimensions[5].height = 26

section(ws, 9, "", 9)
calc(ws["A9"], f'="¿CUÁNTO TENDRÁS A LOS "&{PP["jub"]}&"?"', bold=True, size=11, color=NAVY)
labs = [("A", "PATRIMONIO LÍQUIDO"), ("C", "LO QUE PONES DE TU BOLSILLO"), ("E", "LO QUE GENERAN TUS INVERSIONES"),
        ("G", "EN € DE HOY"), ("I", "INGRESO MENSUAL (€ DE HOY)")]
for c, l in labs:
    ws[f"{c}10"] = l; ws[f"{c}10"].font = font(size=8, bold=True, color="666666")
calc(ws["A11"], f"={atj('G')}", EUR0, bold=True, size=14)
calc(ws["C11"], f"={atj('I')}+{atj('J')}", EUR0, bold=True, size=14, color="1565C0")
calc(ws["E11"], f"={atj('K')}", EUR0, bold=True, size=14, color=GOOD)
calc(ws["G11"], f"={atj('H')}", EUR0, bold=True, size=14)
calc(ws["I11"], f"={atj('H')}*{PP['ret']}/12+{PP['pension']}*14/12", EUR0, bold=True, size=14)
calc(ws["A12"], f'="Por cada euro que pones, tus inversiones añaden "&TEXT(E11/C11,"0.00")&" €. El ingreso mensual suma la regla del "&TEXT({PP["ret"]},"0%")&" y una pensión pública de "&TEXT({PP["pension"]},"#,##0")&" € × 14 pagas (de hoy), repartida en 12 meses."', size=8, color="666666")

# Jubilación
JS = 14
section(ws, JS, "🏖️ TU JUBILACIÓN · ¿hasta cuándo te dura el dinero?", 9)
for c, l in (("A", "SIN DECISIONES"), ("C", "CON TUS DECISIONES"), ("E", "LO QUE SACAS EL 1.er AÑO"), ("G", "PENSIÓN + RENTAS EL 1.er AÑO")):
    ws[f"{c}{JS+1}"] = l; ws[f"{c}{JS+1}"].font = font(size=8, bold=True, color="666666")
def dura(x): return f'=IF({x}>={PP["jub"]}+{DURA_MAX},"Más allá de los "&({PP["jub"]}+{DURA_MAX}),"Hasta los "&{x})'
calc(ws[f"A{JS+2}"], dura(DURA_B), bold=True, size=14, color=NAVY)
calc(ws[f"C{JS+2}"], dura(DURA_C), bold=True, size=14, color=NAVY)
calc(ws[f"E{JS+2}"], f"={PR}$H${J1}", EUR0, bold=True, size=14)
calc(ws[f"G{JS+2}"], f"={PR}$F${J1}+{PR}$G${J1}", EUR0, bold=True, size=14)
calc(ws[f"A{JS+3}"], f'="Cada año sacas lo que no cubren la pensión y las rentas; el capital renta un "&TEXT({PP["r_jub"]},"0%")&" y la pensión sube con la inflación. Con decisiones incluye la cuota si la hipoteca sigue viva."', size=8, color="666666")

# Hitos
HT = JS + 5
section(ws, HT, "🏆 TUS HITOS · se desbloquean solos según tu plan", 9)
header(ws, HT + 1, ["Hito", "Meta", "Hoy", "%", "Año", "Edad", "Faltan (años)", "Estado", "Progreso"])
hitos = [
    ("🛟 Colchón de emergencia", f"={PP['gastos']}*{PP['colchon']}", f"={S0['ef0']}", "D", EUR0),
    ("💶 100.000 € de patrimonio líquido", 100000, f"={S0['liq0']}", "G", EUR0),
    ("💰 150.000 € de patrimonio líquido", 150000, f"={S0['liq0']}", "G", EUR0),
    ("🏠 Vivienda comprada", 1, f"={PI}{g['parte']}", "COMPRA", PCT),
    ("🔓 Libertad financiera 25 %", 0.25, f"={LIB}", "M", PCT),
    ("💎 200.000 € de patrimonio líquido", 200000, f"={S0['liq0']}", "G", EUR0),
    ("🔓 Libertad financiera 50 %", 0.5, f"={LIB}", "M", PCT),
    ("🏖️ Que el dinero te dure hasta los 90", 90, f"={DURA_B}", "DURA", "0"),
]
for i, (name, meta, hoy, c, fmt) in enumerate(hitos):
    r = HT + 2 + i
    ws.cell(row=r, column=1, value=name).font = font()
    if isinstance(meta, str): calc(ws.cell(row=r, column=2), meta, fmt)
    else: inp(ws.cell(row=r, column=2), meta, fmt)
    link(ws.cell(row=r, column=3), hoy, fmt)
    calc(ws.cell(row=r, column=4), f"=MIN(1,IFERROR(C{r}/B{r},0))", PCT)
    if c == "COMPRA":
        flag, tcomp = DD["compra"], DD["tc"]
        calc(ws.cell(row=r, column=5), f'=IF(C{r}>=B{r},"Hoy",IF({flag}="Sí",{PP["anio"]}+{tcomp},"—"))', "0")
        calc(ws.cell(row=r, column=6), f'=IF(C{r}>=B{r},"",IF({flag}="Sí",{PP["edad"]}+{tcomp},""))', "0")
        calc(ws.cell(row=r, column=7), f'=IF(C{r}>=B{r},"",IF({flag}="Sí",{tcomp},""))', "0")
        calc(ws.cell(row=r, column=8), f'=IF(C{r}>=B{r},"✅ Conseguido",IF({flag}="Sí","📅 En DECISIONES","🤔 Por decidir"))')
    elif c == "DURA":
        for cc in (5, 6, 7): calc(ws.cell(row=r, column=cc), '=""')
        calc(ws.cell(row=r, column=8), f'=IF(C{r}>=B{r},"✅ Cubierto","⚠️ Ajusta el plan")')
    else:
        k = f'COUNTIF({col(c)},"<"&B{r})'
        out = f"OR(C{r}>=B{r},{k}>{PR}{TJ})"
        calc(ws.cell(row=r, column=5), f'=IF(C{r}>=B{r},"Hoy",IF({k}>{PR}{TJ},"—",INDEX({col("C")},{k}+1)))', "0")
        calc(ws.cell(row=r, column=6), f'=IF({out},"",INDEX({col("B")},{k}+1))', "0")
        calc(ws.cell(row=r, column=7), f'=IF({out},"",{k})', "0")
        calc(ws.cell(row=r, column=8), f'=IF(C{r}>=B{r},"✅ Conseguido",IF({k}>{PR}{TJ},"⛔ Antes de jubilarte no","🔄 En camino"))')
    calc(ws.cell(row=r, column=9), f'=REPT("█",ROUND(D{r}*10,0))&REPT("░",10-ROUND(D{r}*10,0))&" "&TEXT(D{r},"0%")', color=GOOD)
HE = HT + 1 + len(hitos)

# Logros
LG = HE + 2
section(ws, LG, "🏅 TUS LOGROS DE AHORRO · medalla por % de ingresos ahorrado", 9)
header(ws, LG + 1, ["Año", "Ingresos", "Gastos", "Ahorro", "% ahorrado", "Medalla", "Meses con datos", "", ""])
for i in range(4):
    r = LG + 2 + i
    calc(ws.cell(row=r, column=1), f"={PP['anio']}-1+{i}", "0", bold=True)
    calc(ws.cell(row=r, column=2), f"=SUMIFS({DR('G')},{DR('L')},A{r})", EUR)
    calc(ws.cell(row=r, column=3), f"=SUMIFS({DR('H')},{DR('L')},A{r})", EUR)
    calc(ws.cell(row=r, column=4), f"=B{r}-C{r}", EUR)
    calc(ws.cell(row=r, column=7), f"=COUNTIF({DR('L')},A{r})", "0")
    calc(ws.cell(row=r, column=5), f'=IF(G{r}=0,"",D{r}/B{r})', PCT)
    calc(ws.cell(row=r, column=6), f'=IF(G{r}=0,"",IF(E{r}>=0.4,"🥇 Oro",IF(E{r}>=0.25,"🥈 Plata",IF(E{r}>=0.1,"🥉 Bronce","—")))&IF(G{r}<12," (parcial)",""))')
ws.cell(row=LG + 6, column=1, value="Oro ≥ 40 % · Plata ≥ 25 % · Bronce ≥ 10 %").font = font(size=8, color="666666")

# Decisiones summary
DS = LG + 8
section(ws, DS, "🧭 TUS DECISIONES · ¿qué pasa si…?", 9)
for c, l in (("B", "SIN DECISIONES"), ("D", "CON TUS DECISIONES"), ("G", "DIFERENCIA")):
    ws[f"{c}{DS+1}"] = l; ws[f"{c}{DS+1}"].font = font(size=8, bold=True, color="666666")
calc(ws[f"A{DS+2}"], f'="Patrimonio líquido a los "&{PP["jub"]}', size=9)
calc(ws[f"B{DS+2}"], f"={atj('G')}", EUR0, bold=True, size=12)
calc(ws[f"D{DS+2}"], f"={atj('O')}", EUR0, bold=True, size=12)
calc(ws[f"G{DS+2}"], f"=D{DS+2}-B{DS+2}", EUR, bold=True, size=12)
ws[f"A{DS+3}"] = "Con tus inmuebles (netos de hipoteca)"; ws[f"A{DS+3}"].font = font(size=9)
calc(ws[f"B{DS+3}"], f"={atj('T')}", EUR0, bold=True, size=12)
calc(ws[f"D{DS+3}"], f"={atj('U')}", EUR0, bold=True, size=12)
calc(ws[f"G{DS+3}"], f"=D{DS+3}-B{DS+3}", EUR, bold=True, size=12)
ws[f"A{DS+4}"] = "Libertad financiera"; ws[f"A{DS+4}"].font = font(size=9)
calc(ws[f"B{DS+4}"], f"={atj('M')}", PCT, bold=True, size=12)
calc(ws[f"D{DS+4}"], f"={atj('Q')}", PCT, bold=True, size=12)
calc(ws[f"A{DS+5}"], f'="Decisiones: "&TEXT({DD["lump"]},"#,##0")&" € del efectivo a inversión, +"&TEXT({DD["extra"]},"#,##0")&" €/mes a inversión, −"&TEXT({DD["rec"]},"#,##0")&" €/mes de gastos"&IF({DD["compra"]}="Sí",", compra de la vivienda ("&{DD["esc"]}&")","")&". Comprar pasa dinero líquido a ladrillo: mira las dos filas."', size=8, color="666666")

# Charts
CH = DS + 7
section(ws, CH, "📊 TU FUTURO EN GRÁFICOS", 9)
last = R0 + (JUB - PER["edad"])  # de la edad de hoy a la jubilación
cats = Reference(ws_pro, min_col=2, min_row=R0, max_row=last)
def ref(c): return Reference(ws_pro, min_col=c, min_row=H0, max_row=last)
c1 = AreaChart(); c1.grouping = "stacked"; c1.title = "Tu patrimonio líquido por edad"
for c in (4, 5, 6): c1.add_data(ref(c), titles_from_data=True)
c1.set_categories(cats)
c2 = LineChart(); c2.title = "Patrimonio con inmuebles: sin vs con tus decisiones"
for c in (20, 21): c2.add_data(ref(c), titles_from_data=True)
c2.set_categories(cats)
c3 = AreaChart(); c3.grouping = "stacked"; c3.title = "De dónde sale tu dinero"
for c in (9, 10, 11): c3.add_data(ref(c), titles_from_data=True)
c3.set_categories(cats)
c4 = BarChart(); c4.title = "Libertad financiera: % de tus gastos cubiertos"
c4.add_data(ref(13), titles_from_data=True); c4.set_categories(cats); c4.y_axis.number_format = '0%'
c4.legend = None
c5 = LineChart(); c5.title = "Tu capital en la jubilación: sin vs con decisiones"
c5.add_data(Reference(ws_pro, min_col=9, min_row=JH, max_row=J1 + 30), titles_from_data=True)
c5.add_data(Reference(ws_pro, min_col=14, min_row=JH, max_row=J1 + 30), titles_from_data=True)
c5.set_categories(Reference(ws_pro, min_col=2, min_row=J1, max_row=J1 + 30))
for chx in (c1, c2, c3, c5):
    chx.y_axis.number_format = '#,##0'
for chx in (c1, c2, c3, c4, c5):
    chx.height = 7.5; chx.width = 14; chx.x_axis.title = "Edad"
c5.width = 28.5
ws.add_chart(c1, f"A{CH+1}"); ws.add_chart(c2, f"F{CH+1}")
ws.add_chart(c3, f"A{CH+17}"); ws.add_chart(c4, f"F{CH+17}")
ws.add_chart(c5, f"A{CH+33}")
widths(ws, {"A": 34, "B": 13, "C": 13, "D": 13, "E": 13, "F": 17, "G": 13, "H": 19, "I": 20})

# ================================================================ PANEL_DATOS (lo lee el dashboard)
wsp = wb.create_sheet("PANEL_DATOS")
F_ = "'MI FUTURO'!"
def put(ws_, r, c, v, fmt=None):
    cell = ws_.cell(row=r, column=c, value=v); cell.font = font(size=9)
    if fmt: cell.number_format = fmt
    return cell
row = 1
def tabla(nombre, cabeceras):
    global row
    put(wsp, row, 1, "## " + nombre).font = font(size=10, bold=True, color=NAVY)
    for i, h in enumerate(cabeceras): put(wsp, row + 1, 1 + i, h).font = font(size=9, bold=True)
    row += 2
put(wsp, row, 1, "Esta pestaña la lee el dashboard. No la toques: todo son fórmulas.").font = font(size=9, italic=True, color="888888"); row += 2
tabla("KPI", ["clave", "valor"])
kpis = [
    ("fecha_datos", f"={S0['mes']}", "dd/mm/yyyy"), ("patrimonio_liquido", f"={S0['liq0']}", EUR0),
    ("inmuebles", f"={S0['inm0']}", EUR0), ("patrimonio_total", f"={S0['tot0']}", EUR0),
    ("cuenta_1", f"=INDEX({DR('B')},{NDAT})", EUR0), ("cuenta_2", f"=INDEX({DR('C')},{NDAT})", EUR0), ("cuenta_3", f"=INDEX({DR('D')},{NDAT})", EUR0),
    ("nombre_cuenta_1", CU[0], None), ("nombre_cuenta_2", CU[1], None), ("nombre_cuenta_3", CU[2], None),
    ("libertad", f"={LIB}", PCT), ("nivel", f"={F_}$A$5", None), ("siguiente_nivel", f"={F_}$A$7", None),
    ("ingresos_medio", f"=IFERROR(AVERAGEIFS({DR('G')},{DR('N')},1),0)", EUR0), ("gastos_medio", f"={PP['gastos']}", EUR0),
    ("ahorro_medio", f"={PP['ap_ef']}", EUR0), ("casa_media", f"={CASA_MEDIA}", EUR0),
    ("por_revisar", f"=COUNTIFS({MVR('H')},\"Por revisar\",{MVR('B')},\">=\"&'MIS DATOS'!$A$6)", "0"),
    ("edad", f"={PP['edad']}", "0"), ("edad_jubilacion", f"={PP['jub']}", "0"), ("pension_mes", f"={PP['pension']}", EUR0),
    ("rentabilidad_efectivo", f"={PP['r_ef']}", "0.00%"),
    ("patrimonio_jub", f"={atj('G')}", EUR0), ("patrimonio_jub_hoy", f"={atj('H')}", EUR0), ("ingreso_mes_jub", f"={F_}$I$11", EUR0),
    ("dura_sin", f"={DURA_B}", "0"), ("dura_con", f"={DURA_C}", "0"), ("dura_max", f"={PP['jub']}+{DURA_MAX}", "0"),
    ("patrimonio_jub_con", f"={atj('O')}", EUR0), ("libertad_jub", f"={atj('M')}", PCT), ("libertad_jub_con", f"={atj('Q')}", PCT),
    ("viajes_aportacion", f"={V}$B$4", EUR0), ("viajes_anio", f"={V}$B$7", "0"), ("viajes_saldo_fin_anio", f"={V}$B$10", EUR0),
    ("viajes_aviso", f"={V}$B$12", None),
    ("decisiones_texto", f"={F_}$A${DS+5}", None),
]
for k, f, fmt in kpis:
    put(wsp, row, 1, k); calc(wsp.cell(row=row, column=2), f, fmt, size=9); row += 1
row += 1
grupos = []
for c_, t_, g_ in CFG["categorias"]:
    if t_ == "Gasto" and g_ not in grupos: grupos.append(g_)
tabla("MESES", ["mes", "cuenta_1", "cuenta_2", "cuenta_3", "ingresos", "gastos", "ahorro", "completo"] + ["g_" + g_ for g_ in grupos])
for i in range(60):
    r = 6 + i; R = row + i
    src = lambda c: f"'MIS DATOS'!{c}{r}"
    calc(wsp.cell(row=R, column=1), f'=IF({src("A")}="","",{src("A")})', "yyyy-mm-dd", size=9)
    for j, c in enumerate("BCDGHIN"):
        calc(wsp.cell(row=R, column=2 + j), f'=IF($A{R}="","",{src(c)})', None, size=9)
    for j, g_ in enumerate(grupos):
        extra = {"Comida": f'+SUMIFS({FXR("B")},{FXR("F")},$A{R})', "Transporte": f'+SUMIFS({FXR("C")},{FXR("F")},$A{R})'}.get(g_, "")
        calc(wsp.cell(row=R, column=9 + j), f'=IF($A{R}="","",-SUMIFS({MVR("E")},{MVR("B")},$A{R},{MVR("I")},"{g_}",{MVR("H")},"Gasto"){extra})', None, size=9)
row += 61
tabla("CASA", ["concepto", "media", "total12", "ultimo"])
for i in range(len(casa_cats)):
    r = C0 + i
    for j, c in enumerate("AONP"):
        calc(wsp.cell(row=row, column=1 + j), f"=CASA!{c}{r}", None, size=9)
    row += 1
row += 1
tabla("CASA_MESES", ["mes"] + casa_cats)
for j in range(12):
    colL = ws_casa.cell(row=4, column=2 + j).column_letter
    calc(wsp.cell(row=row, column=1), f"=CASA!{colL}4", "yyyy-mm-dd", size=9)
    for i in range(len(casa_cats)):
        calc(wsp.cell(row=row, column=2 + i), f"=CASA!{colL}{C0+i}", None, size=9)
    row += 1
row += 1
tabla("PROYECCION", ["edad", "liquido", "liquido_con", "con_inmuebles", "con_inmuebles_con", "libertad", "libertad_con"])
for t in range(31):
    r = R0 + t
    for j, c in enumerate("BGOTUMQ"):
        calc(wsp.cell(row=row, column=1 + j), f"='PROYECCIÓN'!{c}{r}", None, size=9)
    row += 1
row += 1
tabla("JUBILACION", ["edad", "capital", "capital_con", "pension_rentas", "gastos"])
for k in range(31):
    r = J1 + k
    for j, c in enumerate("BINFE"):
        f = f"='PROYECCIÓN'!{c}{r}" if c != "F" else f"='PROYECCIÓN'!F{r}+'PROYECCIÓN'!G{r}"
        calc(wsp.cell(row=row, column=1 + j), f, None, size=9)
    row += 1
row += 1
tabla("HITOS", ["hito", "meta", "hoy", "pct", "anio", "estado", "formato"])
for i, h in enumerate(hitos):
    r = HT + 2 + i
    for j, c in enumerate("ABCDEH"):
        calc(wsp.cell(row=row, column=1 + j), f"={F_}{c}{r}", None, size=9)
    put(wsp, row, 7, {EUR0: "eur", PCT: "pct"}.get(h[4], "num"))
    row += 1
row += 1
tabla("VIAJES", ["viaje", "fecha", "tu_parte", "pagado_con", "estado", "saldo_tras", "llega"])
for i in range(VT1 - VT0 + 1):
    r = VT0 + i
    for j, c in enumerate("ABIFGJK"):
        calc(wsp.cell(row=row, column=1 + j), f'=IF(VIAJES!$A{r}="","",VIAJES!{c}{r})', "yyyy-mm-dd" if c == "B" else None, size=9)
    row += 1
row += 1
tabla("VIAJES_ESCENARIOS", ["concepto", "ahora", "a", "b", "c"])
for i in range(9):
    r = 5 + i
    for j, c in enumerate("EHIJK"):
        calc(wsp.cell(row=row, column=1 + j), f"=VIAJES!{c}{r}", None, size=9)
    row += 1
widths(wsp, {"A": 26, "B": 16})
wsp.sheet_state = "visible"

# ================================================================ LEEME
ws = ws_lee
title(ws, "PLAN FINANCIERO PERSONAL · cómo se usa", "Viene con datos ficticios de ejemplo. Para usarlo con los tuyos, sigue docs/GUIA_DESPLIEGUE.md del repositorio.", 2)
lines = [
    ("Cada mes", ""),
    ("1. Descarga", "Los movimientos de tus bancos, en el formato que te da cada uno. Ya se reconocen Openbank (cuenta y tarjeta), CaixaBank e ING; otro banco se añade en BANCOS, en el importador."),
    ("2. Déjalos", "En la carpeta «Movimientos», junto a este sheet (se crea sola la primera vez que importas)."),
    ("3. Importa", "Menú Plan financiero → Importar movimientos. No duplica lo que ya estaba, aunque las descargas se solapen."),
    ("4. Revisa", "En MOVIMIENTOS, lo amarillo (Bizum, ingresos en efectivo…): elige categoría y explica en Nota."),
    ("", ""),
    ("Qué tocas", ""),
    ("MOVIMIENTOS", "Solo las columnas Categoría y Nota."),
    ("REGLAS", "Para que un concepto se categorice solo la próxima vez."),
    ("CATEGORÍAS", "Qué cuenta como ingreso, gasto o traspaso, y qué es de la casa."),
    ("PARÁMETROS", "Edad, jubilación, rentabilidades, pensión. Gastos y ahorro salen de tus datos, pero puedes escribir encima."),
    ("FLEXIBLE", "Si tienes retribución flexible: lo que se descuenta de tu nómina cada mes."),
    ("DECISIONES", "¿Y si invierto parte del efectivo, recorto gastos, compro vivienda ahora o dentro de unos años…?"),
    ("HISTORIAL", "Cada vez que cambias una cifra del plan, apunta aquí qué, de qué a qué y por qué."),
    ("VIVIENDA · INMUEBLES", "Precio, hipoteca y alquiler; y lo que tengas en ladrillo."),
    ("VIAJES", "Tu cuenta de viajes: lo que aportas al mes, tu parte de cada viaje, si llega, y qué pasa si aportas otra cantidad."),
    ("", ""),
    ("Qué miras", ""),
    ("MI FUTURO", "El panel: nivel, patrimonio, jubilación, hitos, medallas de ahorro, decisiones y gráficos."),
    ("CASA", "Gastos de la vivienda por concepto y mes."),
    ("MIS DATOS", "Saldos de fin de mes e ingresos y gastos, solos."),
    ("PANEL_DATOS", "Lo que lee el dashboard. No se toca."),
    ("", ""),
    ("Colores", "Amarillo con texto azul = lo cambias tú · Verde = viene de otra pestaña o de tus datos · Negro = fórmula."),
    ("Ojo", "Es una herramienta de planificación, no asesoramiento financiero ni fiscal. Las rentabilidades son supuestos."),
]
for i, (a, b) in enumerate(lines):
    r = 4 + i
    ws.cell(row=r, column=1, value=a).font = font(bold=True, color=NAVY if b == "" else "000000")
    ws.cell(row=r, column=2, value=b).font = font()
widths(ws, {"A": 20, "B": 110})

# ================================================================ HISTORIAL
ws = ws_his
title(ws, "HISTORIAL · lo que he ido cambiando", "Cada vez que cambias una cifra del plan (aportación a viajes, pensión, una decisión…), apúntalo aquí: la más reciente arriba.", 6)
header(ws, 3, ["Fecha", "Qué cambia", "Antes", "Ahora", "Por qué", "Dónde"])
dvh = DataValidation(type="list", formula1='"' + ",".join(w.title for w in wb.worksheets if w.title != "PANEL_DATOS") + '"', allow_blank=True)
ws.add_data_validation(dvh)
for i in range(200):
    r = 4 + i
    fila = CFG["historial"][i] if i < len(CFG["historial"]) else [None] * 6
    inp(ws.cell(row=r, column=1), d_(fila[0]), "dd/mm/yyyy")
    for j in range(1, 6):
        inp(ws.cell(row=r, column=1 + j), fila[j])
        ws.cell(row=r, column=1 + j).alignment = Alignment(wrap_text=True, vertical="top")
    dvh.add(f"F{r}")
widths(ws, {"A": 12, "B": 24, "C": 24, "D": 32, "E": 48, "F": 15})
ws.freeze_panes = "A4"

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.active = 1
SALIDA = AQUI.parent / "plantilla" / "Plan_financiero_demo.xlsx"
SALIDA.parent.mkdir(exist_ok=True)
wb.save(SALIDA)
print("ok →", SALIDA)
